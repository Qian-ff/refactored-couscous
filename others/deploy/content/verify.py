"""Real Docker/private-protocol acceptance using synthetic files; never parses user files on host."""
from __future__ import annotations

import argparse
import asyncio
from dataclasses import replace
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "others/backend"))
from app.core.content_runner import DockerParseRunner  # noqa: E402
from app.domains.content.parse_runner import ParseTask, RunnerUnavailable  # noqa: E402


def fixtures(root):
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    (root / "safe.txt").write_text("实验安全\n个人防护", encoding="utf-8")
    doc = Document()
    doc.add_paragraph("实验安全")
    doc.save(root / "safe.docx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "实验安全"
    deck.save(root / "safe.pptx")
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    contents = DecodedStreamObject()
    contents.set_data(b"BT /F1 12 Tf 20 200 Td (Lab safety) Tj ET")
    page[NameObject("/Contents")] = contents
    with (root / "safe.pdf").open("wb") as output:
        writer.write(output)
    return {kind: (root / ("safe." + kind)).read_bytes() for kind in ("pdf", "pptx", "docx", "txt")}


async def source(data):
    for offset in range(0, len(data), 65536):
        yield data[offset:offset + 65536]


def task(data, kind="txt"):
    return ParseTask(uuid4(), uuid4(), 1, uuid4(), hashlib.sha256(data).hexdigest(), kind, "utf-8")


async def denied(awaitable):
    try:
        await awaitable
    except RunnerUnavailable:
        return
    raise AssertionError("Private runner accepted an operation that must be fenced")


async def completed(runner, execution):
    for _attempt in range(150):
        current = await runner.inspect(execution)
        if current.state != "running":
            return current
        await asyncio.sleep(.1)
    raise AssertionError("Synthetic parse did not finish")


async def docker(*arguments):
    # Docker inspect/exec output here is fixed trusted metadata/probes, not document output.
    process = await asyncio.create_subprocess_exec("docker", *arguments, stdout=asyncio.subprocess.PIPE,
                                                    stderr=asyncio.subprocess.PIPE)
    stdout, stderr = await asyncio.wait_for(process.communicate(), 15)
    assert len(stdout) <= 65536 and len(stderr) <= 65536
    assert process.returncode == 0, "Trusted Docker verification command failed"
    return stdout


async def verify(args):
    runtime = ROOT / "others/runtime/c02-public"
    config = json.loads((runtime / "control.json").read_text())
    client_config = json.loads((runtime / "parser-client.json").read_text())
    runner, second = DockerParseRunner(client_config), DockerParseRunner(client_config)
    report = {"image": config["image"], "four_formats": {}, "sandbox_verified": False}
    args.report.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=runtime) as temporary:
        for kind, data in fixtures(Path(temporary)).items():
            print("verify: Docker format " + kind, flush=True)
            specification = task(data, kind)
            try:
                await runner.start(specification, source(data))
                state = await completed(runner, specification.execution_id)
                assert state.state == "succeeded" and state.stop_confirmed, state.error_code
                output = await runner.result(specification.execution_id)
                parsed = json.loads(output)
                assert parsed["format"] == kind and parsed["chunks"]
                assert any("Lab safety" in item["text"] if kind == "pdf" else "实验安全" in item["text"]
                           for item in parsed["chunks"])
                accepted = json.loads((Path(config["root"]) / "tasks" / str(specification.execution_id) / "accepted").read_text())
                assert accepted == {"kernel_seals": 15, "input_bytes": len(data),
                                    "input_sha256": specification.input_sha256}
                replay = await runner.start(specification, source(data))
                assert replay.state == "succeeded" and await runner.result(specification.execution_id) == output
                # A different binding must not stop or overwrite the accepted execution.
                await denied(second.start(replace(specification, generation=2), source(data)))
                assert (await runner.inspect(specification.execution_id)).state == "succeeded"
                report["four_formats"][kind] = {"input_bytes": len(data), "output_bytes": len(output),
                    "chunks": len(parsed["chunks"]), "kernel_seals": accepted["kernel_seals"],
                    "identical_replay": True, "different_binding_rejected": True}
            finally:
                assert (await runner.stop(specification.execution_id)).stop_confirmed
                assert await runner.release(specification.execution_id)
                assert await runner.release(specification.execution_id)
                assert not (Path(config["root"]) / "tasks" / str(specification.execution_id)).exists()

    data = b"Synthetic held input"
    specification = task(data)
    waiting, resume = asyncio.Event(), asyncio.Event()

    async def held_source():
        waiting.set()
        yield data
        await resume.wait()

    pending = asyncio.create_task(runner.start(specification, held_source()))
    print("verify: actual container policy and competing clients", flush=True)
    await asyncio.wait_for(waiting.wait(), 15)
    container_name = "labconnect-parse-" + config["instance"] + "-" + specification.execution_id.hex
    try:
        inspected = json.loads(await docker("inspect", "--type", "container", "--format", "{{json .}}", container_name))
        limits, settings = inspected["HostConfig"], inspected["Config"]
        assert inspected["State"]["Running"] and inspected["Image"] == config["image"]
        assert settings["User"] == "65532:65532" and limits["ReadonlyRootfs"]
        assert limits["NetworkMode"] == "none" and limits["CapDrop"] == ["ALL"]
        assert "no-new-privileges:true" in limits["SecurityOpt"]
        assert limits["Memory"] == limits["MemorySwap"] == 384 * 1024 * 1024
        assert limits["NanoCpus"] == 1000000000 and limits["PidsLimit"] == 64
        assert limits["LogConfig"]["Type"] == "none"
        assert "size=46137344" in limits["Tmpfs"]["/work"]
        assert len(inspected["Mounts"]) == 1 and inspected["Mounts"][0]["Destination"] == "/output"
        actual = json.loads(await docker("exec", container_name, "python", "-c", r'''
import errno,json,os,socket
result={"uid":os.getuid(),"socket_absent":not os.path.exists('/var/run/docker.sock'),
        "no_database_or_ai_secrets":not any('SECRET' in key or 'DATABASE' in key or 'API_KEY' in key for key in os.environ)}
try:
    open('/parser/c02-forbidden','wb').close()
    result['root_write_denied']=False
except OSError as error:
    result['root_write_denied']=error.errno in (errno.EROFS,errno.EACCES)
try:
    socket.create_connection(('1.1.1.1',53),timeout=.3).close()
    result['network_denied']=False
except OSError:
    result['network_denied']=True
print(json.dumps(result))
'''))
        assert actual == {"uid": 65532, "socket_absent": True, "no_database_or_ai_secrets": True,
                          "root_write_denied": True, "network_denied": True}
        await denied(second.start(task(b"Second parser"), source(b"Second parser")))
        assert (await runner.inspect(specification.execution_id)).state == "running"
        report["inspect_and_real_access"] = {**actual, "memory_swap_bytes": limits["MemorySwap"],
            "pids": limits["PidsLimit"], "cpus": 1, "work_tmpfs_bytes": 44 * 1024 * 1024,
            "sealed_input_max_bytes": 20 * 1024 * 1024, "two_clients_exclusion": True}
        assert (await second.stop(specification.execution_id)).stop_confirmed
        resume.set()
        await denied(pending)
        assert await runner.release(specification.execution_id)
        report["stop_while_receiving"] = True
    finally:
        resume.set()
        if not pending.done():
            pending.cancel()
        await asyncio.gather(pending, return_exceptions=True)
        await runner.stop(specification.execution_id)
        await runner.release(specification.execution_id)

    late = task(b"Delayed start")
    print("verify: durable stop-before-start tombstone", flush=True)
    assert (await runner.stop(late.execution_id)).stop_confirmed
    assert await runner.release(late.execution_id)
    await denied(runner.start(late, source(b"Delayed start")))
    assert await runner.release(late.execution_id)
    report["stop_before_start_and_release_fence"] = True
    if os.name != "nt":
        raise RuntimeError("This operational restart check requires the registered Windows candidate launcher")
    powershell = shutil.which("pwsh")
    assert powershell
    for script in ("Stop-ContentParser.ps1", "Start-ContentParser.ps1"):
        print("verify: " + script, flush=True)
        # Background Windows launchers can inherit PIPE handles after their parent has exited.
        # Files let us wait for the actual launcher exit, without waiting for a service's EOF.
        with (args.report.parent / (script + ".stdout.txt")).open("wb") as stdout, \
                (args.report.parent / (script + ".stderr.txt")).open("wb") as stderr:
            completed_process = await asyncio.to_thread(subprocess.run,
                [powershell, "-NoProfile", "-File", str(ROOT / "others/scripts" / script)],
                cwd=ROOT, stdout=stdout, stderr=stderr, timeout=60)
        assert completed_process.returncode == 0, "Owned parser restart failed"
    print("verify: tombstone after controller restart", flush=True)
    await denied(runner.start(late, source(b"Delayed start")))
    assert await runner.release(late.execution_id)
    report["restart_preserves_late_start_tombstone"] = True
    report["passed_checks"] = ["four_formats", "actual_docker_policy", "actual_access_denials", "sealed_input",
        "two_clients_exclusion", "stop_during_receipt", "same_binding_replay", "mismatched_binding_rejected",
        "release_cleanup", "release_idempotency", "stop_before_start", "restart_tombstone"]
    # Resource stress, forced recovery and HTTP/PG are distinct outstanding acceptance surfaces.
    report["sandbox_verified"] = False
    report["remaining"] = ["resource_stress_and_timeouts", "forced_control_restart", "real_http_pg_workflow"]
    args.report.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", required=True, type=Path)
    asyncio.run(verify(parser.parse_args()))
