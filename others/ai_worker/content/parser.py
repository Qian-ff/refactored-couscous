"""Local text extraction. Untrusted input must run inside A's restricted ParseRunner."""
from importlib.metadata import version
from pathlib import Path
import re

from content.archive import inspect_office
from content.protocol import MAX_CHUNKS, MAX_TEXT_CHARS, ParseError, encode_result, split_text


class Extraction:
    def __init__(self):
        self.chunks = []
        self.warnings = set()
        self.characters = 0

    def add(self, text, position):
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        self.characters += len(text)
        if self.characters > MAX_TEXT_CHARS:
            raise ParseError("RESOURCE_LIMIT")
        self.chunks.extend(split_text(text, position, len(self.chunks) + 1))
        if len(self.chunks) > MAX_CHUNKS:
            raise ParseError("RESOURCE_LIMIT")


def parse_file(path: Path, kind: str, encoding: str = "utf-8") -> bytes:
    if kind not in {"pdf", "pptx", "docx", "txt", "text"}:
        raise ParseError("UNSUPPORTED_TYPE")
    limit = 1024 * 1024 if kind in {"txt", "text"} else 20 * 1024 * 1024
    sealed = False
    if re.fullmatch(r"/proc/self/fd/[0-9]+", str(path)):
        # Only the fixed Linux wrapper's inherited, kernel-sealed input may use a proc fd path.
        try:
            import fcntl
            descriptor = int(path.name)
            required = fcntl.F_SEAL_WRITE | fcntl.F_SEAL_GROW | fcntl.F_SEAL_SHRINK | fcntl.F_SEAL_SEAL
            sealed = fcntl.fcntl(descriptor, fcntl.F_GET_SEALS) & required == required
        except (ImportError, OSError, ValueError):
            pass
    if not path.is_file() or (path.is_symlink() and not sealed) or not 0 < path.stat().st_size <= limit:
        raise ParseError("RESOURCE_LIMIT")
    output = Extraction()
    parser_version = "labconnect.parse/1"
    try:
        if kind in {"txt", "text"}:
            raw = path.read_bytes()
            if len(raw) > limit:
                raise ParseError("RESOURCE_LIMIT")
            if encoding not in {"utf-8", "gb18030"}:
                raise ParseError("UNSUPPORTED_ENCODING")
            codec = encoding
            if encoding == "utf-8":
                codec = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            text = raw.decode(codec, errors="strict")
            if "\x00" in text:
                raise ParseError("CORRUPT_FILE")
            paragraphs = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
            if len(paragraphs) > 20000:
                raise ParseError("RESOURCE_LIMIT")
            for number, paragraph in enumerate(paragraphs, 1):
                output.add(paragraph, number)
        elif kind == "pdf":
            with path.open("rb") as stream:
                if stream.read(5) != b"%PDF-":
                    raise ParseError("UNSUPPORTED_TYPE")
            from pypdf import PdfReader
            parser_version += ";pypdf=" + version("pypdf")
            reader = PdfReader(path, strict=True)
            if reader.is_encrypted:
                raise ParseError("ENCRYPTED_FILE")
            if len(reader.pages) > 500:
                raise ParseError("RESOURCE_LIMIT")
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if not text.strip():
                    output.warnings.add("PAGES_WITHOUT_TEXT")
                if "/XObject" in page.get("/Resources", {}):
                    output.warnings.add("IMAGE_OR_GRAPHICS_CONTENT")
                if page.get("/Annots"):
                    output.warnings.add("ANNOTATIONS_NOT_EXTRACTED")
                output.add(text, number)
        else:
            names = inspect_office(path, kind)
            parser_version += ";lxml=" + version("lxml")
            if kind == "docx":
                from docx import Document
                from docx.table import Table
                from docx.text.paragraph import Paragraph
                parser_version += ";python-docx=" + version("python-docx")
                document = Document(path)
                position = 0

                def body(container):
                    nonlocal position
                    for item in container.iter_inner_content():
                        if isinstance(item, Paragraph):
                            position += 1
                            if position > 20000:
                                raise ParseError("RESOURCE_LIMIT")
                            output.add(item.text, position)
                        elif isinstance(item, Table):
                            seen = set()
                            for row in item.rows:
                                for cell in row.cells:
                                    if cell._tc not in seen:
                                        seen.add(cell._tc)
                                        body(cell)
                body(document)
                # These structures are not represented by the body iterator, even if empty.
                output.warnings.add("DOCX_NON_BODY_CONTENT_NOT_EXTRACTED")
                if any("comments" in name or "footnotes" in name or "endnotes" in name for name in names):
                    output.warnings.add("DOCX_NOTES_OR_COMMENTS")
            else:
                from pptx import Presentation
                parser_version += ";python-pptx=" + version("python-pptx")
                presentation = Presentation(path)
                if len(presentation.slides) > 300:
                    raise ParseError("RESOURCE_LIMIT")

                def shape_text(shapes):
                    for shape in shapes:
                        if hasattr(shape, "shapes"):
                            yield from shape_text(shape.shapes)
                        elif shape.has_text_frame:
                            yield shape.text_frame.text
                        elif shape.has_table:
                            for row in shape.table.rows:
                                yield "\t".join(cell.text for cell in row.cells)
                        else:
                            output.warnings.add("PPTX_NON_TEXT_OBJECTS")
                for number, slide in enumerate(presentation.slides, 1):
                    if slide._element.get("show") in {"0", "false"}:
                        output.warnings.add("HIDDEN_SLIDES_NOT_EXTRACTED")
                        continue
                    text = "\n".join(shape_text(slide.shapes))
                    if not text.strip():
                        output.warnings.add("SLIDES_WITHOUT_TEXT")
                    output.add(text, number)
                if any(name.startswith("ppt/notesSlides/") for name in names):
                    output.warnings.add("PPTX_NOTES_NOT_EXTRACTED")
        return encode_result(kind, output.chunks, sorted(output.warnings), parser_version)
    except ParseError:
        raise
    except ImportError as error:
        raise ParseError("PARSER_UNAVAILABLE") from error
    except UnicodeError as error:
        raise ParseError("INVALID_ENCODING") from error
    except (OSError, ValueError, KeyError, TypeError, EOFError) as error:
        raise ParseError("CORRUPT_FILE") from error
