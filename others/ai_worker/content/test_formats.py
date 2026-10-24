"""Synthetic real CLI fixtures; dependencies are installed only by A's public package."""
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

READY = all(importlib.util.find_spec(name) is not None for name in ("pypdf", "docx", "pptx", "lxml"))


class FourFormatTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def cli(self, path, kind):
        output, errors = self.root / "output.json", self.root / "errors.txt"
        environment = {key: value for key, value in os.environ.items()
                       if key.upper() in {"SYSTEMROOT", "WINDIR", "PATH", "TEMP", "TMP", "LANG", "LC_ALL"}}
        with output.open("wb") as stdout, errors.open("wb") as stderr:
            completed = subprocess.run([sys.executable, "-m", "content.cli", "--input", str(path), "--format", kind],
                                       cwd=Path(__file__).parent.parent, env=environment, stdin=subprocess.DEVNULL,
                                       stdout=stdout, stderr=stderr, timeout=20, check=False)
        self.assertLessEqual(output.stat().st_size, 16 * 1024 * 1024)
        self.assertLessEqual(errors.stat().st_size, 64 * 1024)
        self.assertEqual(completed.returncode, 0, "trusted synthetic CLI failed")
        result = json.loads(output.read_bytes())
        self.assertEqual(result["protocol"], "labconnect.parse/1")
        self.assertEqual(result["format"], kind)
        return result

    def test_txt_real_cli_unicode(self):
        path = self.root / "safe.txt"
        path.write_text("实验安全\n个人防护", encoding="utf-8")
        result = self.cli(path, "txt")
        self.assertEqual([chunk["text"] for chunk in result["chunks"]], ["实验安全", "个人防护"])

    @unittest.skipUnless(READY, "A has not published the four parser dependencies")
    def test_docx_paragraph_table_body_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph("实验安全")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "个人防护"
        table.cell(0, 1).text = "设备用电"
        document.add_paragraph("应急处置")
        path = self.root / "safe.docx"
        document.save(path)
        result = self.cli(path, "docx")
        self.assertEqual([chunk["text"] for chunk in result["chunks"]], ["实验安全", "个人防护", "设备用电", "应急处置"])
        self.assertEqual([chunk["locator"]["position"] for chunk in result["chunks"]], [1, 2, 3, 4])

    @unittest.skipUnless(READY, "A has not published the four parser dependencies")
    def test_pptx_groups_tables_hidden_and_notes(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "实验安全"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
        table.cell(0, 0).text, table.cell(0, 1).text = "个人防护", "设备用电"
        slide.notes_slide.notes_text_frame.text = "PRIVATE_NOTE"
        hidden = presentation.slides.add_slide(presentation.slide_layouts[6])
        hidden._element.set("show", "0")
        hidden.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "HIDDEN_TEXT"
        path = self.root / "safe.pptx"
        presentation.save(path)
        result = self.cli(path, "pptx")
        text = "\n".join(chunk["text"] for chunk in result["chunks"])
        self.assertIn("实验安全", text)
        self.assertIn("个人防护\t设备用电", text)
        self.assertNotIn("PRIVATE_NOTE", text)
        self.assertNotIn("HIDDEN_TEXT", text)
        self.assertIn("PPTX_NOTES_NOT_EXTRACTED", result["warnings"])
        self.assertIn("HIDDEN_SLIDES_NOT_EXTRACTED", result["warnings"])

    @unittest.skipUnless(READY, "A has not published the four parser dependencies")
    def test_pdf_original_page_and_blank_warning(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        contents = DecodedStreamObject()
        contents.set_data(b"BT /F1 12 Tf 20 200 Td (Lab safety) Tj ET")
        page[NameObject("/Contents")] = contents
        writer.add_blank_page(width=300, height=300)
        path = self.root / "safe.pdf"
        with path.open("wb") as stream:
            writer.write(stream)
        result = self.cli(path, "pdf")
        self.assertIn("Lab safety", result["chunks"][0]["text"])
        self.assertEqual(result["chunks"][0]["locator"]["position"], 1)
        self.assertIn("PAGES_WITHOUT_TEXT", result["warnings"])


if __name__ == "__main__":
    unittest.main()
