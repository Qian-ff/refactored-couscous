"""Bounded, DB-free template codec. The trusted runner still supplies process isolation."""
import csv
import hashlib
import io
import json
import re
import zipfile
from typing import Literal
from xml.parsers import expat

from pydantic import Field, ValidationError, model_validator

from app.domains.content.questions.schemas import DraftContent, StrictModel, canonical_bytes

VERSION = "labconnect.question-transfer/1"
COLUMNS = ("kind", "stem", "options_json", "answer_ids_json", "explanation", "knowledge_points_json", "difficulty", "provenance_json")
INPUT_LIMITS = {"csv": 1048576, "xlsx": 5242880}
OUTPUT_LIMIT = 16777216
MAX_QUESTIONS = 1000
MAX_CELL = 32768
SHEET = "Questions"
SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


class TransferError(ValueError):
    """Safe machine code; never include uploaded text, filenames or parser tracebacks."""


class PreviewRow(StrictModel):
    row: int = Field(ge=3, le=1002, strict=True)
    content: DraftContent | None = None
    error: Literal["column_count", "cell_size_or_control", "text_escape", "formula_prefix", "invalid_question"] | None = None
    duplicate_of: int | None = Field(default=None, ge=3, le=1002, strict=True)


class Preview(StrictModel):
    protocol: Literal["labconnect.question-transfer/1"] = VERSION
    input_hash: str = Field(pattern=r"^[0-9a-f]{64}$")
    rows: tuple[PreviewRow, ...] = Field(min_length=1, max_length=1000)

    @model_validator(mode="after")
    def valid_rows(self):
        seen = {}
        for number, row in enumerate(self.rows, start=3):
            if row.row != number or (row.content is None) == (row.error is None):
                raise ValueError("预检行号或结果无效")
            expected_duplicate = None
            if row.content is not None:
                if row.content.sources:
                    raise ValueError("交换模板不能提供资料权限引用")
                first = seen.setdefault(row.content.payload.duplicate_hash(), number)
                expected_duplicate = first if first != number else None
            if row.duplicate_of != expected_duplicate:
                raise ValueError("预检重复行不一致")
        if len(canonical_bytes(self)) > OUTPUT_LIMIT:
            raise ValueError("预检结果过大")
        return self

    def digest(self):
        return hashlib.sha256(canonical_bytes(self)).hexdigest()


def _check_input(raw, format):
    if format not in INPUT_LIMITS:
        raise TransferError("unsupported_format")
    if not raw or len(raw) > INPUT_LIMITS[format]:
        raise TransferError("input_size")


def _csv_rows(raw):
    try:
        text = raw.decode("utf-8-sig", errors="strict")
        # csv.field_size_limit is process-global. The process default and byte cap apply first;
        # our explicit per-cell bound below is tighter and never changes another API request.
        yield from csv.reader(io.StringIO(text, newline=""), strict=True)
    except (UnicodeError, csv.Error):
        raise TransferError("invalid_csv") from None


def validate_xlsx_archive(raw):
    """Read actual decompressed bytes before openpyxl, including ignored metadata XML."""
    _check_input(raw, "xlsx")
    allowed = re.compile(r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core)\.xml|"
                         r"xl/(?:workbook|styles|sharedStrings)\.xml|xl/_rels/workbook\.xml\.rels|"
                         r"xl/theme/theme1\.xml|xl/worksheets/sheet1\.xml)\Z")
    total = nodes = cells = rows = characters = 0
    names = set()
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if not 5 <= len(entries) <= 16:
                raise TransferError("xlsx_entries")
            for entry in entries:
                if entry.orig_filename != entry.filename or not allowed.fullmatch(entry.filename) or entry.filename in names:
                    raise TransferError("xlsx_part")
                names.add(entry.filename)
                if entry.flag_bits & 1 or entry.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
                    raise TransferError("xlsx_encrypted_or_compression")
                if entry.file_size > OUTPUT_LIMIT or total + entry.file_size > OUTPUT_LIMIT:
                    raise TransferError("xlsx_expanded_size")
                parser = expat.ParserCreate(namespace_separator=" ")
                depth = 0

                def forbidden(*_args):
                    raise TransferError("xlsx_active_xml")

                def start(name, attrs):
                    nonlocal depth, nodes, cells, rows
                    depth += 1
                    nodes += 1
                    if depth > 32 or nodes > 100000 or len(attrs) > 32:
                        raise TransferError("xlsx_xml_size")
                    if name == SPREADSHEET_NS + " f":
                        raise TransferError("xlsx_formula")
                    if name == REL_NS + " Relationship" and attrs.get("TargetMode", "").lower() == "external":
                        raise TransferError("xlsx_external_link")
                    if name == SPREADSHEET_NS + " row":
                        rows += 1
                        row_number = attrs.get("r", "")
                        if (rows > MAX_QUESTIONS + 2 or not re.fullmatch(r"[1-9][0-9]{0,3}", row_number)
                                or int(row_number) > MAX_QUESTIONS + 2):
                            raise TransferError("row_limit")
                    if name == SPREADSHEET_NS + " c":
                        cells += 1
                        if cells > (MAX_QUESTIONS + 2) * len(COLUMNS):
                            raise TransferError("xlsx_cell_limit")
                        ref = attrs.get("r", "")
                        if not re.fullmatch(r"[A-H][1-9][0-9]{0,3}", ref) or int(ref[1:]) > 1002:
                            raise TransferError("xlsx_cell_range")

                def end(_name):
                    nonlocal depth
                    depth -= 1

                def characters_read(value):
                    nonlocal characters
                    characters += len(value)
                    if characters > OUTPUT_LIMIT:
                        raise TransferError("xlsx_text_size")

                parser.StartElementHandler = start
                parser.EndElementHandler = end
                parser.CharacterDataHandler = characters_read
                parser.StartDoctypeDeclHandler = forbidden
                parser.EntityDeclHandler = forbidden
                parser.ExternalEntityRefHandler = forbidden
                parser.ProcessingInstructionHandler = forbidden
                with archive.open(entry) as source:
                    while chunk := source.read(65536):
                        total += len(chunk)
                        if total > OUTPUT_LIMIT:
                            raise TransferError("xlsx_expanded_size")
                        parser.Parse(chunk, False)
                    parser.Parse(b"", True)
            if not {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/worksheets/sheet1.xml"} <= names:
                raise TransferError("xlsx_required_part")
    except (zipfile.BadZipFile, expat.ExpatError, RuntimeError, NotImplementedError, EOFError):
        raise TransferError("invalid_xlsx") from None


def _openpyxl():
    try:
        import openpyxl
        from openpyxl.xml import DEFUSEDXML
    except ImportError:
        raise TransferError("xlsx_dependency_unavailable") from None
    if openpyxl.__version__ != "3.1.5" or not DEFUSEDXML:
        raise TransferError("xlsx_dependency_unavailable")
    return openpyxl


def _xlsx_rows(raw):
    validate_xlsx_archive(raw)
    api = _openpyxl()
    workbook = None
    try:
        workbook = api.load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        if workbook.sheetnames != [SHEET]:
            raise TransferError("xlsx_worksheet")
        sheet = workbook[SHEET]
        sheet.reset_dimensions()
        for row in sheet.iter_rows():
            if len(row) > len(COLUMNS):
                raise TransferError("column_limit")
            if any(cell.value is not None and (not isinstance(cell.value, str) or cell.data_type == "f") for cell in row):
                raise TransferError("xlsx_text_cells_required")
            values = [cell.value if cell.value is not None else "" for cell in row]
            yield values
    except (ValueError, KeyError, TypeError, OSError):
        raise TransferError("invalid_xlsx") from None
    finally:
        if workbook is not None:
            workbook.close()


def _plain_cell(cell):
    if len(cell) > MAX_CELL or any(ord(char) < 32 and char not in "\n\t" for char in cell):
        raise TransferError("cell_size_or_control")
    return cell


def parse_template(raw: bytes, format: str) -> Preview:
    _check_input(raw, format)
    rows = iter(_csv_rows(raw) if format == "csv" else _xlsx_rows(raw))
    marker = next(rows, None)
    if marker is None or len(marker) > len(COLUMNS) or marker[:2] not in (
        [VERSION, "literal-text-v1"], [VERSION, "escaped-text-v1"]
    ) or any(marker[2:]):
        raise TransferError("template_version")
    escaped = marker[1] == "escaped-text-v1"
    if tuple(next(rows, ())) != COLUMNS:
        raise TransferError("template_columns")
    normalized, seen, size = [], {}, 0
    for number, cells in enumerate(rows, start=3):
        if number > MAX_QUESTIONS + 2:
            raise TransferError("row_limit")
        if len(cells) != len(COLUMNS):
            value = PreviewRow(row=number, error="column_count")
        else:
            try:
                if escaped:
                    if any(not cell.startswith("'") or len(cell) > MAX_CELL + 1 for cell in cells):
                        raise TransferError("text_escape")
                    cells = [cell[1:] for cell in cells]
                cells = [_plain_cell(cell) for cell in cells]
                if not escaped and any(cell.lstrip().startswith(("=", "+", "-", "@")) for cell in cells):
                    raise TransferError("formula_prefix")
                kind, stem, options, answers, explanation, knowledge, difficulty, provenance = cells
                content = DraftContent.model_validate({"payload": {"kind": kind, "stem": stem,
                    "options": json.loads(options), "answer_ids": json.loads(answers), "explanation": explanation,
                    "knowledge_points": json.loads(knowledge), "difficulty": difficulty}, "provenance": json.loads(provenance)})
                duplicate = seen.setdefault(content.payload.duplicate_hash(), number)
                value = PreviewRow(row=number, content=content, duplicate_of=duplicate if duplicate != number else None)
            except TransferError as error:
                value = PreviewRow(row=number, error=str(error))
            except (ValidationError, ValueError, RecursionError):
                value = PreviewRow(row=number, error="invalid_question")
        size += len(canonical_bytes(value)) + 1
        if size > OUTPUT_LIMIT - 1024:
            raise TransferError("preview_size")
        normalized.append(value)
    if not normalized:
        raise TransferError("empty_questions")
    return Preview(input_hash=hashlib.sha256(raw).hexdigest(), rows=tuple(normalized))


def _export_rows(contents):
    yield (VERSION, "escaped-text-v1", "", "", "", "", "", "")
    yield COLUMNS
    for number, content in enumerate(contents, start=1):
        if number > MAX_QUESTIONS:
            raise TransferError("row_limit")
        content = DraftContent.model_validate(content)
        if content.sources:
            # Exported locators are not portable grants to another school's materials.
            raise TransferError("export_requires_projected_content")
        payload = content.payload
        values = (payload.kind, payload.stem, canonical_bytes([o.model_dump() for o in payload.options]).decode(),
                  canonical_bytes(payload.answer_ids).decode(), payload.explanation,
                  canonical_bytes(payload.knowledge_points).decode(), payload.difficulty,
                  canonical_bytes(content.provenance).decode())
        yield tuple("'" + _plain_cell(value) for value in values)


def export_template(contents, format: str) -> bytes:
    if format not in INPUT_LIMITS:
        raise TransferError("unsupported_format")
    if format == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, lineterminator="\r\n")
        for row in _export_rows(contents):
            writer.writerow(row)
            if stream.tell() > INPUT_LIMITS[format]:
                raise TransferError("export_size")
        result = stream.getvalue().encode("utf-8-sig")
    else:
        api = _openpyxl()
        workbook = api.Workbook(write_only=True)
        sheet = workbook.create_sheet(SHEET)
        try:
            for row in _export_rows(contents):
                sheet.append(row)
            output = io.BytesIO()
            workbook.save(output)
            result = output.getvalue()
        finally:
            workbook.close()
    # Keep every exported template re-importable under the same input limits.
    if len(result) > INPUT_LIMITS[format]:
        raise TransferError("export_size")
    return result
